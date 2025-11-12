import re
from contextlib import suppress
from decimal import Decimal
from html import escape
from io import BytesIO

import django
import pandas as pd
from django.conf import settings
from django.db.models import Sum
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from weasyprint import HTML

# Setup Django
django.setup()

from src.financials.models import CashBox, Invoice, PersonalAccount, Template
from src.services.models import PaymentDetail


def get_invoice_data(invoice_id):
    invoice = Invoice.objects.select_related(
        "personal_account__apartment__house", "tariff"
    ).get(id=invoice_id)

    payment_detail = PaymentDetail.objects.first()

    owner_name = (
        invoice.personal_account.user.get_full_name()
        if invoice.personal_account and invoice.personal_account.user
        else ""
    )
    house_address = (
        invoice.personal_account.apartment.house.address
        if invoice.personal_account and invoice.personal_account.apartment
        else ""
    )
    apartment_num = (
        invoice.personal_account.apartment.apartment_number
        if invoice.personal_account and invoice.personal_account.apartment
        else ""
    )
    full_address_line = ", ".join(
        part
        for part in [owner_name, house_address, f"кв. {apartment_num}" if apartment_num else ""]
        if part
    )

    data = {
        "Inv_Number": invoice.invoice_number,
        "Status": invoice.get_status_display(),
        "Inv_Date": invoice.date.strftime("%d.%m.%Y") if invoice.date else "",
        "Inv_Period": f"{invoice.start_date:%d.%m.%Y} - {invoice.end_date:%d.%m.%Y}"
        if invoice.start_date and invoice.end_date else "",
        "Проведена": "Проведена" if invoice.conducted else "Не проведена",
        "Владелец": owner_name,
        "Дом": house_address,
        "Квартира": apartment_num,
        "full_address_line": full_address_line,
        "Client_Acc": invoice.personal_account.account_number if invoice.personal_account else "",
        "Телефон": getattr(invoice.personal_account.user, "phone", "") if invoice.personal_account and invoice.personal_account.user else "",
        "Секция": (
            invoice.personal_account.apartment.section.name
            if invoice.personal_account and invoice.personal_account.apartment and invoice.personal_account.apartment.section
            else ""
        ),
        "Payee_Name": payment_detail.description if payment_detail else "",
        "Тариф": invoice.tariff.title if invoice.tariff else "",
    }

    total = Decimal("0")
    for idx, item in enumerate(
        invoice.items.select_related("tariff_service__service", "tariff_service__unit"), start=1
    ):
        service = item.tariff_service
        row = {
            f"Услуга_{idx}": service.service.name if service and service.service else "",
            f"Количество_{idx}": item.count,
            f"Тариф_{idx}": service.tariff.title if service and service.tariff else "",
            f"Ед._изм_{idx}": service.unit.name if service and service.unit else "",
            f"Цена_{idx}": service.price if service else "",
            f"Сумма_{idx}": item.total,
        }
        total += item.total
        data.update(row)

    personal_account_id = invoice.personal_account.id if invoice.personal_account else None
    invoices_total = payments_total = Decimal("0")
    if personal_account_id:
        invoices = Invoice.objects.filter(personal_account_id=personal_account_id, conducted=True).prefetch_related("items")
        invoices_total = sum(
            sum(item.total for item in inv.items.all()) for inv in invoices if inv.status in ("zero", "counted")
        )
        payments_total = CashBox.objects.filter(personal_account_id=personal_account_id, is_conducted=True).aggregate(total=Sum("suma"))["total"] or Decimal("0")

    data.update(
        {
            "invoices_total": invoices_total,
            "Оплачено_всего": payments_total,
            "Client_Balance": payments_total - invoices_total,
            "Total_Amount": total,
        }
    )
    return data


def fill_invoice_to_excel(invoice_id):
    """Generate invoice Excel from DB template in memory."""
    data = get_invoice_data(invoice_id)

    # 🔹 Берем дефолтный шаблон из БД
    template_obj = Template.objects.filter(is_default=True).first()
    if not template_obj:
        raise FileNotFoundError("Default template not found in database")

    file_bytes = template_obj.file.read()
    template_io = BytesIO(file_bytes)
    wb = load_workbook(template_io, keep_vba=True)
    ws = wb.active

    # 🔹 Замена плейсхолдеров
    placeholder_pattern = re.compile(r"\{\{([^}]+)\}\}")
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                new_val = cell.value
                for match in placeholder_pattern.findall(cell.value):
                    if match in data:
                        new_val = new_val.replace(f"{{{{{match}}}}}", str(data[match]))
                cell.value = new_val

    # 🔹 Поиск строки шаблона услуг
    service_row_idx = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and re.search(r"\{\{Услуга(_\d+)?\}\}", cell.value):
                service_row_idx = cell.row
                break
        if service_row_idx:
            break

    if service_row_idx is None:
        raise ValueError("Template row with {{Услуга}} not found")

    template_row = [cell.value for cell in ws[service_row_idx]]
    merged_ranges = [r for r in ws.merged_cells.ranges if r.min_row == service_row_idx]

    for m in merged_ranges:
        with suppress(KeyError):
            ws.unmerge_cells(str(m))
    ws.delete_rows(service_row_idx, 1)

    thin_border = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
    services = sorted((k for k in data if k.startswith("Услуга_")), key=lambda x: int(x.split("_")[1]))

    for idx, _ in enumerate(services):
        row_num = service_row_idx + idx
        ws.insert_rows(row_num, 1)
        for col, template_val in enumerate(template_row, start=1):
            cell = ws.cell(row=row_num, column=col)
            value = template_val
            if isinstance(template_val, str):
                value = re.sub(
                    r"\{\{([\w\.\-А-Яа-яЁёІіЇїЄє]+)\}\}",
                    lambda m: f"{{{{{m.group(1)}_{idx + 1}}}}}",
                    template_val,
                )
                for k, v in data.items():
                    value = value.replace(f"{{{{{k}}}}}", str(v) if v is not None else "")
            cell.value = value.strip() if isinstance(value, str) else value
            cell.border = thin_border

        for m in merged_ranges:
            ws.merge_cells(start_row=row_num, start_column=m.min_col, end_row=row_num, end_column=m.max_col)

    output_io = BytesIO()
    wb.save(output_io)
    output_io.seek(0)
    return output_io


def excel_to_html_openpyxl(xlsx_io):
    """Convert Excel (BytesIO) to HTML string."""
    xlsx_io.seek(0)
    wb = load_workbook(xlsx_io, read_only=False, data_only=True)
    ws = wb.active

    merged_cells_map = {}
    for merged in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged.bounds
        rowspan, colspan = max_row - min_row + 1, max_col - min_col + 1
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                merged_cells_map[(r, c)] = (
                    (rowspan, colspan) if (r, c) == (min_row, min_col) else None
                )

    rows_html = []
    for r in range(ws.min_row, ws.max_row + 1):
        cells_html = []
        for c in range(ws.min_column, ws.max_column + 1):
            key = (r, c)
            if key in merged_cells_map and merged_cells_map[key] is None:
                continue
            cell = ws.cell(row=r, column=c)
            value = escape(str(cell.value)) if cell.value is not None else ""
            attrs = ""
            if merged_cells_map.get(key):
                rowspan, colspan = merged_cells_map[key]
                if rowspan > 1:
                    attrs += f' rowspan="{rowspan}"'
                if colspan > 1:
                    attrs += f' colspan="{colspan}"'
            style_inline = ""
            try:
                if cell.font and cell.font.bold:
                    style_inline += "font-weight:bold;"
                if style_inline:
                    attrs += f' style="{style_inline}"'
            except Exception:
                pass
            cells_html.append(f"<td{attrs}>{value}</td>")
        rows_html.append("<tr>" + "".join(cells_html) + "</tr>")

    html_table = "<table>\n" + "\n".join(rows_html) + "\n</table>"
    page = f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<style>
  table {{ border-collapse: collapse; width: 100%; }}
  td {{ border: 1px solid #ddd; padding: 6px; vertical-align: top; }}
</style>
</head>
<body>
{html_table}
</body>
</html>"""
    return page


def html_to_pdf(html_string):
    """Convert HTML string to PDF in memory (BytesIO)."""
    pdf_io = BytesIO()
    HTML(string=html_string).write_pdf(pdf_io)
    pdf_io.seek(0)
    return pdf_io
